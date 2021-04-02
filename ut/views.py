from django.views.generic import View,UpdateView,CreateView,ListView,DeleteView
from django.http import HttpResponse, HttpResponseRedirect
from django.urls import reverse,reverse_lazy
from django.shortcuts import render
from .tables import *
from .forms import *
from django.template.context_processors import csrf
import numpy as np
from django_tables2 import SingleTableView
#from .filters import ClassesFilter
from django.db.models import Exists, OuterRef
import json
from django.template.loader import render_to_string
from django.http import JsonResponse
from django.utils.html import strip_tags

from .utclasses import *
from django_tables2 import RequestConfig
from django.contrib.auth.mixins import LoginRequiredMixin

def index(request):
    projects=Projects.objects.all()
    default_reports=Projects.objects.filter(DefaultReport__isnull=False)
    dr = {}
    for r in default_reports:
        dr[r.id]=r.DefaultReport.id #get_reporttable(r.DefaultReport_id)['table']
    context=get_base_context()
    #context={'table_list':am.to_dict('records')}
    context['classes']=Classes.objects.all()
    context['projects']=projects
    context['reports']=Reports.objects.all()
    context['defaultreports']=dr

    return render(request,'ut/index.html',context)

from django.forms import ModelForm

def get_base_context(c={}):
    if len(c)==0:
        context={}
    else:
        context=c
    context['base_classes']=Classes.objects.all()
    context['base_projects']=Projects.objects.all()
    return context

def table(request,tablename):
    tm=am[am.TableName==tablename].iloc[0].Table
    qs=tm.objects.all()
    class t (tables.Table):
        class Meta:
            model=tm
            sortable=True
    table = t(qs)
    a='''\
    class mform(ModelForm): #ModelForm example
        class Meta:
            model=tm
            fields=[f.name for f in tm._meta.fields]

    form=mform()
    '''
    context = get_base_context({'tablename':tablename, 'table':  table , 'Class_id': 0})
    return render(request, 'ut/showtable.html', context)

def showclass(request,Class_id):
    classname=Classes.objects.get(pk=Class_id).Class
    table=Attributes.objects.filter(Class_id=Class_id)
    context = get_base_context({'classname':classname,'table':table})
    return render(request, 'ut/showtable.html', context)

class delete_instance(LoginRequiredMixin,DeleteView):
    model = Instances
    success_url = reverse_lazy('ut:index')
    def get(self,*args,**kwargs):
        self.success_url = reverse_lazy('ut:instances',kwargs={'Class_id':kwargs['Class_id'],})
        return self.post(*args,**kwargs)

class update_instances_table(View):
    def get(self,request,*args,**kwargs):
        Class_id=kwargs['Class_id']
        qs = create_rawquery_from_attributes(Class_id, {})
        data = {}
        data['table'] = render_to_string('ut/asynctable.html',
                                         {'table': mytable(Class_id=Class_id, style='TableLayout', data=qs)},
                                         request=request)
        return JsonResponse(data)

from django.template.context_processors import csrf
from crispy_forms.utils import render_crispy_form
from django.test.client import RequestFactory


class edit_instance_base(View):
    def get(self,request,Modal=True,ReadOnly=False,*args,**kwargs):
        Class_id=kwargs['Class_id']
        Instance_id=kwargs['Instance_id']
        defaults={}
        if Instance_id==0:
            if 'ref_attribute' in request.GET:
                defaults[request.GET['ref_attribute']]=request.GET['ref_value']
        form = InstanceForm(defaults=defaults,Class_id=Class_id, Instance_id=Instance_id,ReadOnly=ReadOnly, validation=False)
        context = get_base_context()
        for a1 in Attributes.objects.filter(Class_id=Class_id, DataType_id=DT_Table):
            refclass_id = a1.Ref_Class.id
            refattr = a1.Ref_Attribute.Attribute
            if Instance_id == 0:
                filter = {refattr: -1}
            else:
                filter = {refattr: Instance_id}

            context['columns'+ str(a1.id)] = create_qs_sql(refclass_id)['columns']

            qs = create_rawquery_from_attributes(refclass_id, masterclassfilter=filter)
            table = mytable(Class_id=refclass_id, style='ShortLayout', data=qs)
            context['table' + str(a1.id) + ''] = table
        context['form'] = form
        context['Modal']=Modal
        context['Class_id']=Class_id
        context['Instance_id']=Instance_id
        if Modal:
            data={}
            data['modalformcontent']=render_to_string('ut/_edit_instance_modal.html',context=context,request=request)
            return JsonResponse(data)
        else:
            return render(request, 'ut/edit_instance.html', context)

    def post(self,request,*args,**kwargs):
        Class_id=kwargs['Class_id']
        Instance_id=kwargs['Instance_id']
        Next_id=request.POST.get('next','0')
        form=InstanceForm(request.POST,Class_id=Class_id,Instance_id=Instance_id,ReadOnly=False,validation=True)
        res = {}
        if form.is_valid():
            try:
                save_instance_byname(Class_id=Class_id,Instance_id=Instance_id,instance=form.cleaned_data,passed_by_name=False)
                res['success'] = True
                if request.POST.get('action')=='savenext':
                    #newrequest = RequestFactory().get('/')
                    #newrequest.GET=request.POST.copy()
                    ctx = {}
                    ctx.update(csrf(request))
                    if not Next_id.isnumeric(): #need to show error if no next item
                        res['form_errors']='Instance has been saved but next instance has not been identified.'
                        res['success']=False
                    else:
                        Next_id = int(Next_id)
                        form = InstanceForm(Class_id=Class_id, Instance_id=Next_id, ReadOnly=False,validation=False)
                        form_html = render_crispy_form(form, context=ctx)
                        res['form_html']=form_html
            except BaseException as e:
                res['success']=False
                res['form_errors']= str(e)
            return JsonResponse(res)
        else:
            ctx = {}
            ctx.update(csrf(request))
            form = InstanceForm(request.POST, Class_id=Class_id, Instance_id=Instance_id, ReadOnly=False,
                                validation=False)
            form_html = render_crispy_form(form, context=ctx)
            return JsonResponse({'success':False,'form_html':form_html})
            #return HttpResponseRedirect(reverse('ut:instances', args=(Class_id,)))


class edit_instance(LoginRequiredMixin,edit_instance_base):
    pass

class view_instance(edit_instance_base):
    def get(self,request,ReadOnly=True,*args,**kwargs):
        return super().get(request,ReadOnly=ReadOnly,*args,**kwargs)

def classestree_view(request):
    return render(request, "ut/classestree.html", {'table': Classes.objects.all()})

def reports_view(request,Project_id=0):
    if Project_id==0:
        table = ReportsTable(Reports.objects.all())
    else:
        vl = Reports.objects.filter(id=Project_id).values_list('Reports_m2m', flat=True)
        table =  ReportsTable(Projects.objects.filter(id__in=vl))
    RequestConfig(request, paginate={"per_page": 50}).configure(table)
    context=get_base_context()
    context['table']= table
    context['Project_id']= Project_id
    return render(request,'ut/reports.html',context)

def projects_view(request):
    table = ProjectsTable(Projects.objects.all())
    RequestConfig(request, paginate={"per_page": 50}).configure(table)
    context=get_base_context()
    context['table']= table
    return render(request,'ut/projects.html',context)

def classes_view(request,Project_id=0):
    if Project_id==0:
        table = ClassesTable(Classes.objects.all())
    else:
        vl = Reports.objects.filter(id=Project_id).values_list('Classes_m2m', flat=True)
        table =  ClassesTable(Classes.objects.filter(id__in=vl))
    RequestConfig(request, paginate={"per_page": 25}).configure(table)
    context=get_base_context({'table':table,'classes':table,'Project_id':Project_id})
    return render(request,'ut/classes.html',context)

class instances_view():
    def __init__(self,Class_id,filter={}):
        super().__init__(self);
        self.Class_id=Class_id
        self.filter=filter
    def get(self,request,Class_id,SaveToExl):
        self.filter = {}
        for key,value in request.GET.items():
            if (value!='')&(key not in ['sort','page']):
                pass

class ReportRun(View):
    template='ut/showreport.html'
    def get(self,request,Report_id,*args,**kwargs):
        r=Reports.objects.get(pk=Report_id)
        sql=r.Query
        cursor=con.cursor()
        cursor.execute(sql)
        t=dictfetchall(cursor)
        extra_columns=[(c[0], tables.Column()) for c in cursor.description]
        context=get_base_context()
        #context['table']=ReportTable(data=t,extra_columns=extra_columns)
        context['Report_id']=Report_id
        context['ReportName']=r.Report
        return render(request, self.template, context)

def instances(request,Class_id,SaveToExl=False,Project_id=0):
    if Class_id == 0 and Project_id != 0:
        Class_id = Projects.objects.filter(id=Project_id).values_list('Classes_m2m').first()[0]

    defaultfiler = {}
    filterform = InstanceFilterForm(Class_id=Class_id, filter=defaultfiler)  # create_filter_set(Class_id,filter)

    if SaveToExl:
        qs = create_rawquery_from_attributes(Class_id, filter)
        return export_instances_xls(request,qs)
    else:
        pass

    context = get_base_context({'tablename':'Instances','Class_id':Class_id, 'table':table,'filterform':filterform})
    #context['columns']=create_qs_sql(Class_id)['columns']

    if Project_id!=0:
        context['Project']=Projects.objects.get(pk=Project_id)
    return render(request, 'ut/showproject.html',  context)

def attributes_view(request,Class_id):
    attlist=Attributes.objects.filter(Class_id=Class_id)
    table = AttributeTable(attlist)
    RequestConfig(request, paginate={"per_page": 20}).configure(table)
    context=get_base_context({'table':table,'Class_id':Class_id})
    return render(request,'ut/attributes.html',context)

class filters_view(View):
    template='ut/filters.html'
    def get(self,request,*args,**kwargs):
        self.Class_id=kwargs['Class_id']
        context=get_base_context()
        context['Class_id']=self.Class_id
        table = FilterTable(Filters.objects.filter(Class_id=self.Class_id))
        context['table']=table
        return render(request,self.template,context)

def edit_attribute(request,Attribute_id):
    at=Attributes.objects.get(pk=Attribute_id)
    form = AttributeForm(instance=at)
    context=get_base_context({'form':form})
    return render(request,'ut/edit_attribute.html',context)

class BaseContext():
    def get_context_data(self, **kwargs):
        context=super(BaseContext,self).get_context_data(**kwargs)
        context.update(get_base_context())
        return context

class ProjectEdit(BaseContext,UpdateView):
    model = Projects
    template_name = 'ut/edit_attribute.html'
    form_class = ProjectForm
    def get_success_url(self):
        return reverse_lazy('ut:projects_view')

class ProjectCreateVeiw(BaseContext,CreateView):
    model = Projects
    template_name = 'ut/edit_attribute.html'
    form_class = ProjectForm
    def get_success_url(self):
        return reverse_lazy('ut:projects_view')

class ReportEdit(BaseContext,UpdateView):
    model = Reports
    template_name = 'ut/edit_attribute.html'
    form_class = ReportForm
    def get_success_url(self):
        return reverse_lazy('ut:reports_view')

class ReportCreateVeiw(BaseContext,CreateView):
    model = Reports
    template_name = 'ut/edit_attribute.html'
    form_class = ReportForm
    def get_success_url(self):
        return reverse_lazy('ut:reports_view')

#from bootstrap_modal_forms.generic import BSModalCreateView

class FilterCreateView(BaseContext,LoginRequiredMixin,CreateView):
    model = Filters
    template_name = 'ut/edit_attribute.html'
    form_class = FilterEditForm

    def get_success_url(self):
        return reverse_lazy('ut:filters_view', args = (self.Class_id,))

    def get_initial(self):
        initial = super().get_initial()
        self.Class_id = self.kwargs['Class_id']
        initial['Class']=Classes.objects.get(pk=self.Class_id)
        return initial

    def get_form_kwargs(self):
        kw=super(FilterCreateView,self).get_form_kwargs()
        kw['initial']['Class_id']=self.Class_id
        return kw

class FilterUpdateView(BaseContext,LoginRequiredMixin,UpdateView):
    model = Filters
    template_name = 'ut/edit_attribute.html'
    form_class = FilterEditForm

    def get_success_url(self):
        return reverse_lazy('ut:filters_view',args=(self.object.Class.id,))

class AttributeCreateView(BaseContext,LoginRequiredMixin,CreateView):
    model = Attributes
    template_name = 'ut/edit_attribute.html'
    form_class = AttributeForm

    def get_success_url(self):
        return reverse_lazy('ut:attributes_view', args = (self.Class_id,))

    def get_initial(self):
        initial = super().get_initial()
        self.Class_id = self.kwargs['Class_id']
        initial['Class']=Classes.objects.get(pk=self.Class_id)
        return initial

    def get_form_kwargs(self):
        kw=super(AttributeCreateView,self).get_form_kwargs()
        kw['initial']['Class_id']=self.Class_id
        return kw

class AttributeUpdateView(BaseContext,LoginRequiredMixin,UpdateView):
    model = Attributes
    template_name = 'ut/edit_attribute.html'
    form_class = AttributeForm

    def get_success_url(self):
        return reverse_lazy('ut:attributes_view',args=(self.object.Class.id,))

class ClassesUpdateView(BaseContext,LoginRequiredMixin,UpdateView):
    model = Classes
    template_name = 'ut/edit_attribute.html'
    form_class = ClassesForm

    def get_success_url(self):
        return reverse_lazy('ut:classes_view')

class ClassesCreateView(BaseContext,LoginRequiredMixin,CreateView):
    model = Classes
    template_name = 'ut/edit_attribute.html'
    form_class = ClassesForm

    def get_success_url(self):
        return reverse_lazy('ut:classes_view')

import xlwt

from django.http import HttpResponse

import csv

def export_instances_csv(request,raw_qs):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="Instances.csv"'
    writer = csv.writer(response)
    writer.writerow(list(raw_qs.columns))
    rows = raw_queryset_as_values_list(raw_qs)
    for row in rows:
        writer.writerow(row)
    return response

from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

def export_instances_xls(request,raw_qs):
    response = HttpResponse(content_type='application/ms-excel')
    response['Content-Disposition'] = 'attachment; filename="instances.xlsx"'
    wb = Workbook()
    ws = wb.active
    ws.title='Instances'
    # Sheet header, first row
    ws.append(raw_qs.columns)
    rows=raw_queryset_as_values_list(raw_qs)
    for r in rows:
        ws.append(r)
    wb.save(response)
    return response

def export_instances_xls_2(request,raw_qs):
    response = HttpResponse(content_type='application/ms-excel')
    response['Content-Disposition'] = 'attachment; filename="instances.xlsx"'
    wb = xlwt.Workbook(encoding='utf-8')
    ws = wb.add_sheet('Instances')
    # Sheet header, first row
    row_num = 0
    font_style = xlwt.XFStyle()
    font_style.font.bold = True
    columns = raw_qs.columns
    for col_num in range(len(columns)):
        ws.write(row_num, col_num, columns[col_num], font_style)
    # Sheet body, remaining rows
    font_style = xlwt.XFStyle()
    rows = raw_queryset_as_values_list(raw_qs)
    for row in rows:
        row_num += 1
        for col_num in range(len(columns)):
            ws.write(row_num, col_num, row[col_num], font_style)
    wb.save(response)
    return response

#from somewhere import handle_uploaded_file

from django.db.models import Count, F
from django.db import transaction
#@transaction.atomic
def handle_uploaded_file(f,Class_id,commitevery=0,errors='raise'):
    ignore_conflicts= (errors=='ignore')
    totalins=pd.read_excel(f)
    totallen=len(totalins)
    ranges=[]
    if commitevery==0:
        ranges.append([0,len(totallen)])
    else:
        rnum= totallen // commitevery
        rres= totallen % commitevery
        for i in range(0,rnum):
            ranges.append([i*commitevery,(i+1)*commitevery])
        if rres > 0:
            ranges.append([rnum*commitevery,rnum*commitevery+rres])

    for irange in ranges:
        print ('saving class',Class_id,'from',irange[0],'to',irange[1])
        with transaction.atomic():
            ins=totalins[irange[0]:irange[1]]
            user=get_current_user()
            ctime=datetime.now()
            if not ('Code' in ins.columns):
                ins['Code']=ins.apply(lambda x: get_next_counter(Class_id))
            newinstances=ins.Code.apply(lambda x: Instances(Class_id=Class_id,Code=x,Updatedby=user,Updated=ctime))
            Instances.objects.bulk_create(list(newinstances),ignore_conflicts=ignore_conflicts)
            ids = pd.DataFrame(list(Instances.objects.filter(Class_id=Class_id,Code__in=list(ins.Code)).values()))
            ins=pd.merge(ins,ids[['id','Code']],on='Code',how='inner')

            #print ([r for r in ins.iterrows()])
            fl=get_editfieldlist(Class_id)
            for attr in fl:
                #att=Attributes.objects.get(pk=rec.id).Attribute
                if attr.Attribute in ins.columns and (attr.Attribute!='Code'):
                    dt=attr.DataType_id
                    fieldname=get_fieldname(dt)
                    ins_name=attr.Attribute
                    if dt in DTG_Instance:
                        if attr.Ref_Attribute_id == 0:
                            attids = pd.DataFrame(list(Instances.objects.filter(Class_id=attr.Ref_Class_id, Code__in=list(ins[attr.Attribute]))).values('id', 'Code'))
                            attids = attids.rename(columns={'Code': attr.Attribute, 'id': attr.Attribute + '_id'})
                            ins=pd.merge(ins, attids, on=[attr.Attribute], how='left')
                        else:
                            attids = pd.DataFrame(list(Values.objects.filter(Attribute_id=attr.Ref_Attribute_id,char_value__in=list(ins[attr.Attribute])).values('Instance_id', 'char_value')))
                            attids = attids.rename(columns={'char_value': attr.Attribute, 'Instance_id': attr.Attribute + '_id'})
                            ins=pd.merge(ins, attids, on=[attr.Attribute], how='left')
                        values=list(ins.apply(lambda x: Values(**{'Instance_id':x.id, 'Attribute_id':attr.id, 'instance_value_id': x[attr.Attribute + '_id']}),axis=1))
                    else:
                        if dt in DTG_Int:
                            ins[attr.Attribute]=ins[attr.Attribute].apply(lambda x: int(x) if pd.notnull(x) else x)
                        elif dt in DTG_Date:
                            ins[attr.Attribute]=ins[attr.Attribute].apply(lambda x: pd.to_datetime(x).replace(tzinfo=pytz.UTC))

                        values=list(ins.apply(lambda x: Values(**{'Instance_id':x.id,'Attribute_id':attr.id,fieldname.replace('"',''):x[attr.Attribute]}),axis=1))
                    Values.objects.bulk_create(values,ignore_conflicts=ignore_conflicts)

def load_instances(request,Class_id=0):
    if request.method == 'POST':
        form = UploadInstances(request.POST, request.FILES)
        if form.is_valid() and request.FILES['file']:
            commitevery=request.POST.get('Commit')
            if commitevery.isnumeric():
                commitevery = int(commitevery)
            else:
                commitevery = 0
            handle_uploaded_file(request.FILES['file'],Class_id,commitevery,errors=request.POST.get('Errors'))
            return HttpResponseRedirect(reverse('ut:instances', args=(Class_id,)))
    else:
        form = UploadInstances()
    context=get_base_context({'form': form,'Class_id':Class_id})
    return render(request, 'ut/loadinstances.html',context)

class ProtectView(LoginRequiredMixin, View) :
    def get(self, request):
        return render(request, 'ut/index.html')

class FormTemplateView(View):
    template = 'ut/formtemplate.html'
    def get(self, request,Class_id):
        table=get_editfieldlist(Class_id)
        context=get_base_context()
        if Layouts.objects.filter(Class=Class_id).exists():
            context['layout']=Layouts.objects.get(Class=Class_id).FormLayout
        context['table']=table # .to_dict('records')
        context['Class_id']=Class_id
        return render(request, self.template, context)

    def post(self, request,Class_id) :
        if request.POST:
            #layout=json.loads(request.POST['layout'])
            lo_to_save=request.POST['layout']
            if Layouts.objects.filter(Class=Class_id).exists():
                rec=Layouts.objects.get(Class=Class_id)
                rec.FormLayout=lo_to_save
            else:
                rec=Layouts(FormLayout=lo_to_save,Class_id=Class_id)
            rec.save()

        return HttpResponseRedirect(reverse('ut:change_formtemplate',kwargs={'Class_id':Class_id}))

class TableTemplateView(View):
    template = 'ut/tabletemplate.html'
    def get(self, request,Style,Class_id):
        table= get_editfieldlist(Class_id).to_dict('records')
        context=get_base_context()
        if Layouts.objects.filter(Class=Class_id).exists():
            context['layout']=getattr(Layouts.objects.get(Class=Class_id),Style)
        context['table']=table
        context['Class_id']=Class_id
        context['Style']=Style
        return render(request, self.template, context)

    def post(self, request,Style,Class_id) :
        if request.POST:
            #layout=json.loads(request.POST['layout'])
            lo_to_save=request.POST['layout']
            if Layouts.objects.filter(Class=Class_id).exists():
                rec=Layouts.objects.get(Class=Class_id)
            else:
                rec=Layouts(Class=Class_id)
            setattr(rec,Style,lo_to_save)
            rec.save()

        return HttpResponseRedirect(reverse('ut:change_tabletemplate',kwargs={'Style':Style,'Class_id':Class_id}))

from django.forms import formset_factory

class TestForm(forms.Form):
    id=forms.IntegerField()
    name=forms.CharField()

from crispy_forms.layout import LayoutObject, Submit, Row, Column, MultiField,Field,Div,Fieldset,TEMPLATE_PACK,HTML
from crispy_forms.helper import FormHelper

class FormSet1(LayoutObject):
    template='ut/test.html'
    def __init__(self,formset_name_in_context,template=None):
        self.formset_name_in_context = formset_name_in_context
        self.fields=[]
        if template:
            self.template=template
        def render (self,form,form_style,context,template_pack=TEMPLATE_PACK):
            formset = context[self.formset_name_in_context]
            return render_to_string(self.template,{'formset':formset})

class NpEncoder(json.JSONEncoder):
    def default(self, obj):
        if isinstance(obj, np.integer):
            return int(obj)
        elif isinstance(obj, np.floating):
            return float(obj)
        elif isinstance(obj, np.ndarray):
            return obj.tolist()
        else:
            return super(NpEncoder, self).default(obj)

def ajax_change_master(request,Attribute_id):
    value = request.GET['value']
    if value.isnumeric():
        value=int(value)
    else:
        value = 0
    attr=get_attribute(Attribute_id)
    Ref_Class_id=attr.Ref_Class_id
    Class_id=attr.Class_id
    data={}
    data['MasterAttribute_id']=Ref_Class_id
    attrs={}
    #find all attributes who depends on that Attribute_id
    for a in Attributes.objects.filter(Class_id=Class_id):
            #df_attributes[(df_attributes.MasterAttribute_id==Attribute_id)&(df_attributes.Class_id==Class_id)].iterrows()
        if a.MasterAttribute_id==Attribute_id:
            instances=get_options(a.id,values={attr.Attribute:value})
            attrs[a.id]=instances
    data['attrs']=attrs

    #find lookups
    lookups={}
    for a in Attributes.objects.filter(Class_id=Class_id,DataType_id=DT_Lookup,InternalAttribute_id=Attribute_id):
            #df_attributes[(df_attributes.DataType_id==DT_Lookup)&(df_attributes.InternalAttribute_id==Attribute_id)].iterrows():
        if value !=0 :
            lookups[a.id]=Values.objects.filter(Attribute_id=a.Ref_Attribute_id,Instance_id=value).values()[0]['char_value']
        else:
            lookups[a.id]='(None)'
    data['lookups']=lookups
    return JsonResponse(data,encoder=NpEncoder)

def ajax_get_class_columns(request,Class_id):
    cqs = Instances.objects.raw(create_qs_sql(Class_id=Class_id)['sql'] + ' limit 0')
    return JsonResponse({'columns':cqs.columns})

from querystring_parser import parser

def ajax_get_report_data (request, Report_id,sample=0, filter={}):
    r = Reports.objects.get(pk=Report_id)
    sql = r.Query
    cursor = con.cursor()
    cursor.execute(sql)
#    desc = cursor.description
    res={}
    res['data']=dictfetchall(cursor)
    res['recordsTotal']=len(res['data'])
    res['recordsFiltered'] = len(res['data'])
    #res['columns']=[col[0] for col in cursor.description]
    return JsonResponse(res)

def ajax_get_report_columns (request, Report_id,filter={}):
    r = Reports.objects.get(pk=Report_id)
    sql = 'select * from (' + r.Query + ') original limit 0'
    cursor = con.cursor()
    cursor.execute(sql)
#    desc = cursor.description
    res={}
    #res['data']=dictfetchall(cursor)
    res['columns']=[col[0] for col in cursor.description]
    return JsonResponse(res)


def ajax_get_attribute_options(request,Class_id,Attribute_id,maxrecords=10):
    validation=False
    values=request.GET
    term=values.get('term')
    if pd.isnull(term):
        term = ''
    attr=get_attribute(Attribute_id)
    instances=[]
    if (attr.MasterAttribute_id>0):
        m_attr=get_attribute(attr.MasterAttribute_id)
        tmp=values.get(m_attr.Attribute)
        if pd.notnull(tmp):
            m_value=int(tmp)
        else:
            m_value=0
        if attr.Ref_Attribute_id == 0:
            for r in Values.objects.filter(instance_value_id=m_value,Instance__Class__id=attr.Ref_Class_id,Instance__Code__icontains=term)[0:maxrecords]:
                instances.append({'id':r.Instance_id,'text':r.Instance.Code})
        else:
            for r in Values.objects.filter(Attribute_id=attr.Ref_Attribute_id,char_value__icontains=term,
                    Instance_id__in=Values.objects.filter(instance_value_id=m_value,Instance__Class_id=attr.Ref_Class_id).values_list('Instance_id',flat=True))[0:maxrecords]:
                instances.append({'id': r.Instance_id, 'text': r.char_value})

                #Instances.objects.raw(select_options_sql.format(val=m_value,cl=attr.Ref_Class_id,att=attr.Ref_Attribute_id) + " and  v2.char_value like '%{}%'".format(term))[0:maxrecords]:
                #instances.append({'id': r.id, 'text': r.char_value})
    else:
        if attr.Ref_Attribute_id == 0:
            for r in Instances.objects.filter(Class_id=attr.Ref_Class_id,Code__icontains=term)[0:maxrecords]:
                instances.append({'id': r.id, 'text': r.Code})
        else:
            for r in Values.objects.filter(Attribute_id=attr.Ref_Attribute_id,char_value__icontains=term)[0:maxrecords]:
                instances.append({'id': r.Instance_id, 'text': r.char_value})
    return JsonResponse({'success':True,'results': instances,'pagination':{'more':True}})

def ajax_get_class_data(request,Class_id):
    res={}
    masterfilter={}
    if 'filtername' in request.GET:
        masterfilter = {request.GET['filtername']:request.GET['filtervalue']}
    filter={}
    if 'filterform' in request.GET:
        if len(request.GET['filterform'])>0:
            filterform=json.loads(request.GET['filterform'])
            for f in filterform:
                filter[f['name']]=f['value']
    offset=None
    limit=None
    search=''
    draw=0
    ssargs = parser.parse(request.META['QUERY_STRING'])
    print ('META.query_string',request.META['QUERY_STRING'])
    if 'start' in request.GET:
        offset=int(ssargs.get('start'))
        limit=int(ssargs.get('length'))
        draw=int(ssargs.get('draw'))
        search=ssargs['search']['value']

    orderby={}
    if ssargs:
        for key,val in ssargs.get('order').items():
            orderby[ssargs['columns'][int(val['column'])]['data']]=val['dir']

    sql=create_rawquery_sql(Class_id=Class_id,filter=filter,masterclassfilter=masterfilter,orderby=orderby,
                                       search=search,offset=offset,limit=limit)

    count_sql=create_count_sql(Class_id=Class_id
                               ,filter=filter,masterclassfilter=masterfilter,search=search
                               )
    with con.cursor() as cursor:
        cursor.execute(count_sql)
        rec=cursor.fetchone()
    recordsTotal= rec[1]

    res['data'] = raw_queryset_as_dict(sql)
    res['recordsTotal']=recordsTotal
    res['recordsFiltered'] = recordsTotal #len(res['data']) #len(res['data'])
    #recordsFiltered
    res['draw']=draw+1
    return JsonResponse(res)

from .sendouts import send_mail

class send_instance_email(View):
    def get(self,request,Class_id,Modal=True,*args,**kwargs):
        Instance_id = request.GET.get('Instance_id')
        if pd.isnull(Instance_id) or Instance_id == 0:
            instance=request.GET
        else:
            instance=get_instance(Class_id,Instance_id)
        #adjust names for the template
        instance_adj={}
        for k,v in instance.items():
            instance_adj[k.replace(' ','_')]=v
        form=SendInstanceEmailForm(Class_id=Class_id,instance=instance_adj)
        context = get_base_context()
        context['form']=form
        context['Class_id']=Class_id
        if Modal:
            data={}
            data['success']=True
            data['modalformcontent']=render_to_string('ut/showclassemail.html',context=context,request=request)
            return JsonResponse(data)
        else:
            return render(request, 'ut/showclassemail.html',context)

    def post(self,request,Class_id,Instance_id=0, Modal=True,*args,**kwargs):
        form=SendInstanceEmailForm(request.POST,Class_id=Class_id,instance={ })
        if form.is_valid():
            email_subject= request.POST['subject']
            html_body = request.POST['text_body']
            plain_body = strip_tags(html_body)
            to_list = [request.POST['to']]
            send_mail(email_subject,plain_body,settings.EMAIL_HOST_USER,to_list,html_message=html_body)
            print ('email sent')
        if Modal:
            return JsonResponse({'success':True,'data':None})
        else:
            return HttpResponseRedirect(reverse('ut:instances', args=(Class_id,)))

class emailtemplates_view(View):
    def get(self,request,*args,**kwargs):
        tlist=EmailTemplates.objects.all()
        table = EmailTemplateTable(tlist)
        RequestConfig(request, paginate={"per_page": 20}).configure(table)
        context=get_base_context({'table':table})
        return render(request,'ut/templates.html',context)

class EmailTemplateCreateView(BaseContext,LoginRequiredMixin,CreateView):
    model = EmailTemplates
    template_name = 'ut/edit_attribute.html'
    form_class = EmailTemplateForm
    def get_success_url(self):
        return reverse_lazy('ut:templates_view')

class EmailTemplateUpdateView(BaseContext,LoginRequiredMixin,UpdateView):
    model = EmailTemplates
    template_name = 'ut/edit_attribute.html'
    form_class = EmailTemplateForm
    def get_success_url(self):
        return reverse_lazy('ut:templates_view')
